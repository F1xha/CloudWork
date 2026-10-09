"""Create an honest report using only evidence files actually present."""
from pathlib import Path
from datetime import datetime, timezone
import json
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'evidence';OUT.mkdir(exist_ok=True)
resultfile=OUT/'resultado.json'
result=json.loads(resultfile.read_text()) if resultfile.exists() else {}
passed=result.get('status')=='PASS'
doc=Document();sec=doc.sections[0]
sec.top_margin=Inches(.7);sec.bottom_margin=Inches(.7);sec.left_margin=Inches(.8);sec.right_margin=Inches(.8)
sec.page_height=Inches(11.7);sec.page_width=Inches(8.3)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2']:
 st=doc.styles[name];st.font.name='Calibri';st.font.color.rgb=RGBColor(0,0,0)
doc.styles['Normal'].font.size=Pt(11)
doc.styles['Normal'].paragraph_format.space_after=Pt(7)
doc.styles['Normal'].paragraph_format.line_spacing=1.08
doc.styles['Title'].font.size=Pt(26)
doc.styles['Heading 1'].font.size=Pt(17)
doc.styles['Heading 2'].font.size=Pt(13)
# Remove inherited decorative title borders and theme colors.
for style in doc.styles:
 for border in list(style.element.iter(qn('w:pBdr'))): border.getparent().remove(border)
header=sec.header.paragraphs[0];header.text='DUOC UC  ·  CLOUDWORK  ·  EV2'
for run in header.runs:
 run.font.size=Pt(9);run.font.color.rgb=RGBColor(0,0,0)
footer=sec.footer.paragraphs[0];footer.alignment=2
footer.add_run('Michael Giraldo  |  ')
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)
def p(t): doc.add_paragraph(t)
def h(t): doc.add_heading(t,1)
def page(): doc.add_page_break()
def code(t):
 par=doc.add_paragraph();par.paragraph_format.space_after=Pt(9)
 r=par.add_run(t);r.font.name='DejaVu Sans Mono';r.font.size=Pt(8.5)
def table(headers,rows,widths):
 t=doc.add_table(rows=1, cols=len(headers));t.autofit=False
 for c,w in zip(t.columns,widths): c.width=Inches(w)
 for i,text in enumerate(headers): t.rows[0].cells[i].text=text
 for row in rows:
  cells=t.add_row().cells
  for i,v in enumerate(row):cells[i].text=str(v)
 for ri,row in enumerate(t.rows):
  for i,c in enumerate(row.cells):
   c.width=Inches(widths[i]);c.vertical_alignment=1
   pr=c._tc.get_or_add_tcPr();sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'24445C' if ri==0 else ('F1F5F8' if ri%2 else 'FFFFFF'));pr.append(sh)
   borders=OxmlElement('w:tcBorders')
   for side in ['top','left','bottom','right']:
    el=OxmlElement('w:'+side);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');borders.append(el)
   pr.append(borders)
   margins=OxmlElement('w:tcMar')
   for side in ['top','left','bottom','right']:
    el=OxmlElement('w:'+side);el.set(qn('w:w'),'90');el.set(qn('w:type'),'dxa');margins.append(el)
   pr.append(margins)
   for par in c.paragraphs:
    par.paragraph_format.space_after=Pt(3);par.paragraph_format.space_before=Pt(3)
    for run in par.runs:run.font.size=Pt(10);run.font.bold=(ri==0);run.font.color.rgb=RGBColor.from_string('FFFFFF' if ri==0 else '000000')
  nr=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(nr)
 t.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
 doc.add_paragraph().paragraph_format.space_after=Pt(1)

doc.add_heading('Comunicación de incidencias con RabbitMQ',0)
p('Diseño e implementación de microservicios · Evaluación Parcial 2')
p('Michael Giraldo\nDuoc UC\n5 de octubre de 2026\nActividad equivalente al 10% de EV2 · Informe sin presentación')
h('Propósito y contexto')
p('Diseñamos una extensión del Mini Help Desk de CloudWork para registrar incidencias de soporte. La creación necesita comprobar que la categoría existe. Una vez guardada la incidencia, la auditoría y una notificación interna a soporte pueden procesarse de forma independiente mediante RabbitMQ.')
p('La decisión principal consiste en validar la categoría por HTTP antes del guardado y publicar el evento IncidenciaCreada para dos consumidores. Así, detener auditoría o notificaciones no debería impedir que el usuario registre una incidencia cuando el broker está disponible.')
h('Alcance de la entrega')
p('El módulo ev2-rabbitmq contiene cuatro procesos Spring Boot, un broker y bases H2 separadas por servicio persistente. Comparte un artefacto Java con perfiles de ejecución para reducir la configuración académica. El backend y el frontend de EV1 se conservan; la demostración EV2 utiliza solicitudes HTTP independientes del frontend y de Azure.')
if passed:
 p('Estado de verificación: el archivo resultado.json registra PASS para las pruebas de integración. Las evidencias y capturas disponibles se incorporan al final de este informe. Los resultados corresponden exclusivamente a esa ejecución.')
else:
 p('Estado de verificación: implementación preparada, pendiente de ejecución. El entorno de preparación no disponía de Docker y no pudo acceder a Maven Central. GitHub rechazó la escritura con HTTP 403. Este documento no acredita funcionamiento ni incluye capturas simuladas; debe completarse con la ejecución del procedimiento incluido.')
h('Responsabilidades')
table(['Servicio','Responsabilidad'],[('Incidencias','Valida por HTTP, guarda y publica el evento.'),('Categorías','Responde si existe la categoría solicitada.'),('Auditoría','Conserva el registro del evento recibido.'),('Notificaciones','Conserva un aviso interno consultable para soporte.')],[1.5,5.1])
page();h('Arquitectura y recorrido de la operación')
doc.add_picture(str(ROOT/'docs/arquitectura.png'),width=Inches(6.6))
p('Figura 1. Flujo diseñado. Las flechas hacia las colas representan bindings con la misma routing key; no son llamadas HTTP a los consumidores.')
for x in [
'1. El cliente envía POST /api/incidencias con título, descripción, categoría y prioridad.',
'2. El servicio valida el payload y consulta GET /categorias/{id}. Una categoría inexistente produce HTTP 400; una dependencia no disponible produce HTTP 503. En ambos casos se interrumpe el flujo antes de guardar.',
'3. Si la categoría existe, se genera un identificador de incidencia y otro de evento. La incidencia se guarda en H2 con publicación pendiente.',
'4. El producer publica IncidenciaCreada en el exchange direct. Ambos bindings coinciden con incidencia.creada y cada cola recibe una copia.',
'5. El producer espera la confirmación del broker, no la respuesta de los consumers, y devuelve HTTP 201. El campo publicacion indica CONFIRMADA o PENDIENTE.',
'6. Cada consumidor procesa su copia y guarda su resultado en una base independiente. El eventoId permite relacionar la solicitud, la publicación y ambos resultados.']:p(x)
page();h('Justificación de las comunicaciones')
table(['Comunicación','Modalidad','Razón'],[
('Incidencias a categorías','Síncrona HTTP','La respuesta determina si se puede aceptar la incidencia.'),
('Evento a auditoría','Asíncrona AMQP','Registrar el historial no modifica la aceptación inicial.'),
('Evento a notificaciones','Asíncrona AMQP','El aviso puede procesarse después sin bloquear al usuario.')],[1.6,1.25,3.75])
p('El ACK del broker confirma la publicación, pero no el trabajo posterior. El producer puede esperar esa confirmación hasta cinco segundos. La prueba de consumidores detenidos permite observar la independencia entre aceptar la solicitud y procesar los efectos posteriores.')
h('Contrato y topología de RabbitMQ')
table(['Elemento','Definición'],[
('Evento y producer','IncidenciaCreada · servicio de incidencias'),
('Exchange','incidencias.exchange · direct · durable'),
('Routing key','incidencia.creada'),
('Cola de auditoría','incidencias.auditoria.queue · durable'),
('Cola de notificaciones','incidencias.notificaciones.queue · durable'),
('Bindings','Ambas colas se vinculan al exchange con incidencia.creada.')],[1.6,5])
p('Usamos dos colas para entregar una copia a cada responsabilidad. Si ambos consumidores compartieran una única cola, competirían por los mensajes. Un exchange direct permite seleccionar el evento por coincidencia exacta de routing key y binding key [1].')
h('Payload mínimo')
code('''{
  "eventoId": "UUID del evento",
  "incidenciaId": "UUID de la incidencia",
  "titulo": "Impresora no imprime",
  "categoriaId": 1,
  "prioridad": "ALTA",
  "fechaCreacion": "2026-10-05T22:00:00Z"
}''')
p('Ejemplo ilustrativo del contrato, no evidencia de ejecución. La fecha se genera en UTC. La descripción completa queda en la base principal; auditoría y notificaciones no la necesitan para este flujo.')
page();h('Implementación y separación del código')
p('Usamos Java 17, Spring Boot 3.5.5, Spring AMQP y JDBC con H2. Docker Compose construye una imagen del módulo y ejecuta perfiles incidencias, categorias, auditoria y notificaciones en contenedores distintos. RabbitMQ utiliza la imagen 4.1.4-management.')
table(['Capa','Archivos principales'],[
('API','IncidentController, CategoryController y ProcessedController'),
('Aplicación','IncidentService, AuditService y NotificationService'),
('Dominio','Contracts con solicitud, respuesta y evento'),
('Infraestructura','CategoryClient, MessagingConfig, EventPublisher, consumers y stores JDBC')],[1.35,5.25])
h('Publicación y enrutamiento')
p('MessagingConfig declara el exchange, las colas y ambos bindings. El fragmento corresponde al código entregado.')
code('''var exchange = new DirectExchange(EXCHANGE, true, false);
var audit = new Queue(AUDIT, true);
var notify = new Queue(NOTIFY, true);
BindingBuilder.bind(audit).to(exchange).with(ROUTING);
BindingBuilder.bind(notify).to(exchange).with(ROUTING);''')
p('EventPublisher usa RabbitTemplate con JSON, entrega persistente, mandatory, publisher returns y confirmaciones correlacionadas. Cada mensaje incluye su eventoId como messageId. Si se produce un retorno o una confirmación negativa, la aplicación no marca la publicación como confirmada.')
h('Procesamiento consumidor')
code('''@RabbitListener(queues = MessagingConfig.AUDIT)
public void receive(IncidenciaCreada event) {
    service.process(event);
    log.info("CONSUMED consumer=auditoria eventoId={} incidenciaId={}",
        event.eventoId(), event.incidenciaId());
}''')
p('El adaptador AMQP delega en el servicio de aplicación; la persistencia se encapsula en ProcessedStore. El listener retorna después del guardado y el contenedor confirma el consumo. El consumidor de notificaciones utiliza su propia cola y base con el mismo patrón.')
h('Notificación implementada')
p('La notificación consiste en un aviso interno persistido y consultable en GET /procesados del puerto 8084. No se conecta a un proveedor de correo ni envía WhatsApp. Auditoría usa el puerto 8083 y conserva un registro con una finalidad distinta.')
page();h('Pruebas y evidencias de implementación')
p('El script scripts/verify.py se conecta a servicios y broker reales. Comprueba códigos HTTP, persistencia y el mismo eventoId en ambos consumidores. Captura RabbitMQ Management durante la acumulación de mensajes y después del consumo.')
status='PASS registrado' if passed else 'Pendiente de ejecución'
table(['Prueba','Resultado exigido','Estado'],[
('Registro válido','201 y evento presente en ambos consumers',status),
('Categoría inexistente','400 y ninguna incidencia nueva',status),
('Payload inválido','400 para un título vacío',status),
('Consumers detenidos','201 confirmado y mensaje Ready en cada cola',status),
('Reinicio de consumers','Ambos procesan y las colas quedan vacías',status),
('Categorías detenido','503 y ninguna incidencia nueva',status)],[1.55,3.6,1.45])
p('Las tres pruebas unitarias de IncidentService verifican el rechazo previo al guardado, el orden validar–guardar–publicar y la respuesta PENDIENTE ante fallo del broker. Son pruebas con mocks y no sustituyen la integración. Sus resultados se conservan en surefire-reports cuando Maven se ejecuta.')
h('Archivos que permiten verificar el recorrido')
p('01-solicitud-y-respuesta.json contiene el POST y su respuesta. 02-auditoria.json y 03-notificaciones.json permiten comparar el eventoId. 05-colas-con-mensajes.json y 07-colas-consumidas.json muestran el estado del broker. 09-exchange.json y 10-bindings.json describen la topología. servicios.log reúne los registros SYNC, PUBLISHED y CONSUMED.')
if passed:
 p('La ejecución registra los siguientes identificadores: evento inicial '+result.get('eventoInicial','')+' y evento diferido '+result.get('eventoDiferido','')+'. Estos valores deben coincidir con los JSON y logs adjuntos.')
else:
 p('Estos nombres identifican los archivos que generará la prueba; no significa que existan ya. En esta versión no hay resultados reales de RabbitMQ. La revisión estática confirmó la estructura del proyecto y la sintaxis de los scripts Python, pero la compilación y el comportamiento en ejecución quedan por verificar.')
h('Criterio para cerrar la actividad')
p('La entrega estará verificada cuando las pruebas finalicen correctamente y el informe incorpore capturas reales del exchange con sus dos bindings, colas con mensajes pendientes y colas tras su consumo. El workflow incluido automatiza esa generación; este documento debe regenerarse y revisarse después de ejecutar las pruebas.')
page();h('Ejecución y limitaciones')
p('Desde ev2-rabbitmq, con Docker Desktop iniciado, se ejecuta docker compose up -d --build. La primera construcción requiere internet. El script scripts/demo.ps1 crea una incidencia y consulta ambos consumidores. RabbitMQ Management se abre en http://localhost:15672 con usuario ev2 y contraseña ev2-local-demo.')
code('''docker compose up -d --build
python -m pip install playwright python-docx
python -m playwright install chromium
python scripts/verify.py
python scripts/build_report.py''')
p('Antes de regenerar el informe se deben recoger los logs con docker compose logs --no-color y guardarlos en evidence/servicios.log. El workflow hace ese paso automáticamente y publica el paquete evidencias-ev2-rabbitmq como artefacto de Actions.')
h('Límites del alcance académico')
p('El catálogo contiene tres categorías fijas. El módulo se ejecuta localmente, sin autenticación de aplicación, con puertos vinculados a 127.0.0.1. Los datos persistentes de cada responsabilidad viven en volúmenes separados. Compartir código y contratos en un artefacto simplifica el ejercicio, aunque acopla su versionado.')
p('El guardado y la publicación no son una transacción distribuida. Si el proceso falla entre ambos pasos, o el broker no confirma, puede quedar una incidencia pendiente. La respuesta expone ese estado y no existe recuperación automática. Un outbox transaccional sería una mejora posterior; no se garantiza entrega exactamente una vez.')
p('El merge por eventoId evita duplicar filas en los consumidores. No implementamos DLQ, clustering ni retries avanzados. Los mensajes rechazados no se reencolan automáticamente, por lo que un error consumidor puede descartar un mensaje. La notificación interna no representa un correo efectivamente enviado.')
h('Conclusión')
p('El diseño reserva la comunicación síncrona para una decisión necesaria antes del registro y usa RabbitMQ para dos efectos posteriores independientes. Las colas separadas permiten procesar auditoría y notificaciones a su propio ritmo. '+('La ejecución registrada confirma los casos incluidos en el script; no elimina las limitaciones descritas.' if passed else 'La implementación entregada debe ejecutarse para confirmar esta conducta y cerrar la evidencia de la Parte 2.'))
h('Referencias')
p('[1] RabbitMQ. Routing con Spring AMQP. https://www.rabbitmq.com/tutorials/tutorial-four-spring-amqp')
p('[2] RabbitMQ. Publish Subscribe con Java. https://www.rabbitmq.com/tutorials/tutorial-three-java')
p('[3] Repositorio base CloudWork. https://github.com/F1xha/CloudWork')
p('Enunciado de la actividad proporcionado por el curso. Fechas solicitadas: diseño 2 de octubre e implementación 5 de octubre de 2026. El documento no afirma una entrega previa del diseño.')
for name,title in [('colas-pendientes','Colas con consumidores detenidos'),('exchange-bindings','Exchange y bindings'),('colas-consumidas','Colas después del procesamiento')]:
 f=OUT/(name+'.png')
 if f.exists():
  page();h(title);p('Captura producida por scripts/verify.py contra RabbitMQ Management durante la ejecución.');doc.add_picture(str(f),width=Inches(6.6))
file=OUT/'Informe_EV2_RabbitMQ.docx';doc.save(file);print(file)
